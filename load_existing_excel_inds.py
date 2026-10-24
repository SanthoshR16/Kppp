import openpyxl, os, json

existing_inds = set()

for fp in [
    r"C:\Users\Admin\Desktop\San (Don't Use)\Tenders New.xlsx",
    r"C:\Users\Admin\Desktop\San (Don't Use)\popupdoc\popupdoc\tenders.xlsx"
]:
    if not os.path.exists(fp):
        continue
    try:
        wb = openpyxl.load_workbook(fp, read_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for row in ws.iter_rows(values_only=True):
                if not row or len(row) < 3:
                    continue
                ind = str(row[2] or "").strip()
                dept = str(row[1] or "").strip()
                if ind and ind != "IND No":
                    existing_inds.add(ind)
                    if dept and dept != "Department":
                        existing_inds.add(f"{dept}/{ind}")
        wb.close()
    except Exception as e:
        print(f"Error {fp}: {e}")

print(f"Total existing unique IND keys loaded from Excel: {len(existing_inds)}")
print("Sample keys:", list(existing_inds)[:10])

with open("data/existing_excel_inds.json", "w") as f:
    json.dump(sorted(list(existing_inds)), f, indent=2)
print("Saved to data/existing_excel_inds.json")
