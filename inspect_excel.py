import openpyxl, glob, os

files = [
    r"C:\Users\Admin\Desktop\San (Don't Use)\popupdoc\popupdoc\tenders.xlsx",
    r"C:\Users\Admin\Desktop\San (Don't Use)\Tenders New.xlsx",
]

for fp in files:
    if os.path.exists(fp):
        print("="*60)
        print(f"File: {fp}")
        try:
            wb = openpyxl.load_workbook(fp, read_only=True)
            print("Sheets:", wb.sheetnames)
            for sname in wb.sheetnames:
                ws = wb[sname]
                rows = list(ws.iter_rows(values_only=True))
                print(f"  Sheet '{sname}': {len(rows)} rows")
                if rows:
                    print("    Header:", rows[0][:8])
                    for r in rows[1:4]:
                        print("    Row:", r[:8])
            wb.close()
        except Exception as e:
            print("  Error reading:", e)
    else:
        print(f"Not found: {fp}")
